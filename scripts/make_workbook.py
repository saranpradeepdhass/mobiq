"""Build one .xlsx with every Mobiq worksheet (headers + optional SAMPLE rows) for upload to Google Drive.

Usage:  python scripts/make_workbook.py [--sample] [--out sheets_template/Mobiq.xlsx]
Requires openpyxl (development helper only; not needed to run the app).
All cells are written as text so phone numbers, dates and ids are not converted by Google Sheets.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from core.models import TABLES  # noqa: E402
from services.demo_data import sample_rows  # noqa: E402


def build(out: Path, sample: bool) -> None:
    rows = sample_rows() if sample else {}
    wb = Workbook()
    wb.remove(wb.active)
    for table, (_, headers) in TABLES.items():
        ws = wb.create_sheet(table)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for r in rows.get(table, []):
            ws.append([str(r.get(h, "")) for h in headers])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.number_format = "@"  # text
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(h) + 4)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--sample", action="store_true", help="include clearly marked SAMPLE rows")
    ap.add_argument("--out", type=Path, default=ROOT / "sheets_template" / "Mobiq.xlsx")
    a = ap.parse_args()
    build(a.out, a.sample)
